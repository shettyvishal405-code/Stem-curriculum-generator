"""
Framework Parser Agent
Handles both Excel (.xlsx) and text-based framework documents.
"""

import json
import openpyxl
from pathlib import Path
from .base_agent import BaseAgent
from utils.prompts import framework_parser_prompt


class FrameworkParserAgent(BaseAgent):
    """Parses a framework document into structured session data."""

    def __init__(self, model: str, max_tokens: int = 4096):
        super().__init__(
            name="FrameworkParser",
            model=model,
            max_tokens=max_tokens,
            temperature=0.0,
        )

    def parse(self, framework_input: str, grade: int = None) -> dict:
        path = Path(framework_input)
        if path.exists() and path.suffix in [".xlsx", ".xls"]:
            return self._parse_excel(str(path), grade)
        else:
            return self._parse_text(framework_input)

    def _parse_excel(self, filepath: str, grade: int) -> dict:
        wb = openpyxl.load_workbook(filepath, data_only=True)

        sheet_name = f"Grade {grade}"
        if sheet_name not in wb.sheetnames:
            for sn in wb.sheetnames:
                if str(grade) in sn:
                    sheet_name = sn
                    break
            else:
                raise ValueError(f"No sheet for Grade {grade}. Available: {wb.sheetnames}")

        ws = wb[sheet_name]
        print(f"  [FrameworkParser] Parsing sheet: '{sheet_name}' ({ws.max_row} rows)")

        topics = []
        current_topic = None
        cumulative_session = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            col_topic = str(row[0]).strip() if row[0] else ""
            col_count = row[1] if len(row) > 1 else None
            col_outcome = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            col_tangible = str(row[3]).strip() if len(row) > 3 and row[3] else ""

            if col_topic and col_topic != "None":
                session_count = int(col_count) if col_count else 0
                start_session = cumulative_session + 1
                cumulative_session += session_count

                current_topic = {
                    "topic": col_topic,
                    "session_count": session_count,
                    "start_session": start_session,
                    "end_session": cumulative_session,
                    "learning_outcomes": [],
                    "tangible_outcome": col_tangible if col_tangible and col_tangible != "None" else "",
                }
                topics.append(current_topic)

            if col_outcome and col_outcome != "None" and current_topic:
                current_topic["learning_outcomes"].append(col_outcome)

        sessions = []
        for topic_block in topics:
            sc = topic_block["session_count"]
            if sc == 0:
                continue
            for i in range(sc):
                session_num = topic_block["start_session"] + i
                session = {
                    "session_number": session_num,
                    "topic": topic_block["topic"],
                    "subtopics": [],
                    "learning_outcomes": topic_block["learning_outcomes"],
                    "tangible_outcome": topic_block["tangible_outcome"],
                    "suggested_activities": [],
                    "tools_software": [],
                    "duration_minutes": 50,
                    "prerequisites": [],
                    "session_position_in_topic": i + 1,
                    "total_sessions_in_topic": sc,
                    "notes": "",
                }
                if sc > 1:
                    if i == 0:
                        session["notes"] = f"First of {sc} sessions. Introduce foundational concepts."
                    elif i == sc - 1:
                        session["notes"] = f"Last of {sc} sessions. Consolidate learning and produce tangible outcome."
                    else:
                        session["notes"] = f"Session {i+1} of {sc}. Build progressively."
                sessions.append(session)

        result = {
            "subject": f"ENpower STEM Lab — Grade {grade}",
            "grade": grade,
            "total_sessions": cumulative_session,
            "target_grades": [grade],
            "board": "CBSE/ICSE",
            "topic_blocks": topics,
            "sessions": sessions,
            "extraction_warnings": [],
            "source_file": filepath,
            "source_sheet": sheet_name,
        }
        self._validate_extraction(result)
        print(f"  [FrameworkParser] Extracted {len(topics)} topic blocks, {len(sessions)} sessions")
        return result

    def _parse_text(self, framework_text: str) -> dict:
        system_prompt = framework_parser_prompt()
        user_message = f"""Parse the following framework document and extract ALL session details.

## FRAMEWORK DOCUMENT
---
{framework_text}
---

Extract every session with its topic, learning outcomes, activities, tools, and tangible outcomes.
Return the result as a JSON object following the specified format."""

        result = self.call(system_prompt, user_message)
        self._validate_extraction(result)
        return result

    def _validate_extraction(self, data: dict):
        warnings = data.get("extraction_warnings", [])
        if "sessions" not in data or not data["sessions"]:
            raise ValueError("Framework parser extracted zero sessions.")
        for session in data["sessions"]:
            sn = session.get("session_number", "?")
            if not session.get("topic"):
                warnings.append(f"Session {sn}: No topic")
            if not session.get("learning_outcomes"):
                warnings.append(f"Session {sn}: No learning outcomes")
        data["extraction_warnings"] = warnings
        if warnings:
            print(f"  [FrameworkParser] {len(warnings)} extraction warnings")
